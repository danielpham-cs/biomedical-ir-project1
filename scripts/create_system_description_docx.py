"""Create the concise English system-description report for Project 1."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "output" / "docx" / "helix_system_description_final_v2.docx"
BLACK = RGBColor(0, 0, 0)
BODY = RGBColor(35, 48, 56)
MUTED = RGBColor(80, 94, 103)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    tc_pr.append(shading)


def remove_title_border(document: Document) -> None:
    """Keep Word's required Title style but remove its default blue rule."""
    style_ppr = document.styles["Title"].element.find(qn("w:pPr"))
    if style_ppr is not None:
        borders = style_ppr.find(qn("w:pBdr"))
        if borders is not None:
            style_ppr.remove(borders)


def set_cell_margins(cell, top: int = 90, start: int = 110, bottom: int = 90, end: int = 110) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def write_body(paragraph, text: str) -> None:
    run = paragraph.add_run(text)
    run.font.name = "Aptos"
    run.font.size = Pt(10.3)
    run.font.color.rgb = BODY
    paragraph.paragraph_format.line_spacing = 1.08
    paragraph.paragraph_format.space_after = Pt(5.5)


def add_heading(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(8)
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run(text)
    run.bold = True
    run.font.name = "Aptos Display"
    run.font.size = Pt(11.8)
    run.font.color.rgb = BLACK


def add_bullet(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.left_indent = Inches(0.22)
    paragraph.paragraph_format.first_line_indent = Inches(-0.14)
    paragraph.paragraph_format.space_after = Pt(2.5)
    paragraph.paragraph_format.line_spacing = 1.04
    run = paragraph.add_run(text)
    run.font.name = "Aptos"
    run.font.size = Pt(10.1)
    run.font.color.rgb = BODY


def add_workflow(document: Document) -> None:
    table = document.add_table(rows=1, cols=3)
    table.autofit = False
    widths = (Inches(2.13), Inches(2.13), Inches(2.13))
    items = (
        ("1  INPUT", "PubMed or PMC XML"),
        ("2  PROCESS", "Rules, tokens and index"),
        ("3  OUTPUT", "Ranked results and statistics"),
    )
    for cell, width, (label, value) in zip(table.rows[0].cells, widths, items):
        cell.width = width
        set_cell_shading(cell, "F2F7F7")
        set_cell_margins(cell)
        label_run = cell.paragraphs[0].add_run(label)
        label_run.font.name = "Aptos"
        label_run.font.size = Pt(8.3)
        label_run.font.bold = True
        label_run.font.color.rgb = RGBColor(0, 116, 108)
        body_paragraph = cell.add_paragraph()
        body_paragraph.paragraph_format.space_before = Pt(2)
        body_run = body_paragraph.add_run(value)
        body_run.font.name = "Aptos"
        body_run.font.size = Pt(9.2)
        body_run.font.color.rgb = BODY
    document.add_paragraph().paragraph_format.space_after = Pt(1)


def build_document() -> None:
    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(0.66)
    section.bottom_margin = Inches(0.66)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)

    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.3)
    normal.font.color.rgb = BODY
    normal.paragraph_format.space_after = Pt(5.5)
    normal.paragraph_format.line_spacing = 1.08
    remove_title_border(document)

    if "Report Subtitle" not in document.styles:
        subtitle = document.styles.add_style("Report Subtitle", WD_STYLE_TYPE.PARAGRAPH)
        subtitle.font.name = "Aptos"
        subtitle.font.size = Pt(10.2)
        subtitle.font.color.rgb = MUTED

    document.core_properties.title = "Helix Biomedical Information Retrieval System"
    document.core_properties.subject = "Biomedical Information Retrieval Project 1"

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.paragraph_format.space_after = Pt(2)
    title_run = title.add_run("Helix Biomedical Information Retrieval System")
    title_run.font.name = "Aptos Display"
    title_run.font.size = Pt(18)
    title_run.font.bold = True
    title_run.font.color.rgb = BLACK

    subtitle = document.add_paragraph(style="Report Subtitle")
    subtitle.paragraph_format.space_after = Pt(11)
    subtitle.add_run("System description for Biomedical Information Retrieval Project 1")

    opening = document.add_paragraph()
    write_body(
        opening,
        "Helix is a small biomedical information retrieval system built for this project. "
        "It reads biomedical XML records, prepares the text with clear rules, searches a local collection, "
        "and shows both ranked results and document statistics. The system is designed to make each step visible. "
        "A user can search a paper by PMID or PMCID, upload XML files, or search the local collection with keywords.",
    )

    add_workflow(document)

    add_heading(document, "1. Data Input and XML Processing")
    paragraph = document.add_paragraph()
    write_body(
        paragraph,
        "The system accepts two standard biomedical XML formats: PubMed XML and PMC JATS XML. "
        "PubMed XML usually provides bibliographic information, an abstract, and keywords. PMC JATS XML may also provide full body text. "
        "For a PubMed record, Helix joins all AbstractText sections in their original order and adds the supplied keywords as a final Keywords line. "
        "For a PMC record, it selects the main article abstract and ignores extra teaser, graphical, or editorial abstracts. "
        "Titles, authors, journal names, dates, and DOI values remain metadata and are excluded from text statistics and matching.",
    )

    add_heading(document, "2. Text Processing and Document Statistics")
    paragraph = document.add_paragraph()
    write_body(
        paragraph,
        "Before analysis, Helix normalizes whitespace by changing repeated spaces, new lines, and tabs into one space. "
        "Its rule-based tokenizer keeps biomedical text readable and consistent. For example, decimal numbers such as 0.65, percentages such as 24.9%, "
        "Greek letters, internal hyphens in SARS-CoV-2, and apostrophes in Alzheimer's are handled carefully. "
        "The statistics scope is the abstract plus the final Keywords line. The system reports characters with spaces, characters without spaces, total words, "
        "unique words, term frequency, sentence count, and average sentence length. Word statistics keep stopwords and do not use stemming, so they remain a direct count of the source text.",
    )

    add_heading(document, "3. Retrieval, Matching and Ranking")
    paragraph = document.add_paragraph()
    write_body(
        paragraph,
        "For retrieval, Helix applies lowercasing, tokenization, English stopword removal when appropriate, and Porter stemming. "
        "Porter stemming groups common suffix variants such as test, tested, and testing. The processed terms are stored in an inverted index, which maps each term to the documents and positions where it occurs. "
        "This allows the system to search efficiently without reading every XML file again for each query.",
    )
    add_bullet(document, "AND mode returns documents that contain every valid query term.")
    add_bullet(document, "OR mode returns documents that contain at least one valid query term.")
    add_bullet(document, "PHRASE mode keeps stopwords and requires the terms to appear consecutively in the same paragraph.")
    paragraph = document.add_paragraph()
    write_body(
        paragraph,
        "After candidate documents are selected, Helix ranks them with BM25. BM25 considers term frequency, how rare a term is in the collection, and document length. "
        "The implementation uses k1 = 1.2 and b = 0.75. A BM25 score is useful for comparing documents returned by the same query; it is not a percentage or an accuracy value.",
    )

    add_heading(document, "4. Rule Based Sentence Segmentation")
    paragraph = document.add_paragraph()
    write_body(
        paragraph,
        "Sentence count is not based on a simple split at every period. Helix uses a rule-based end-of-sentence method. It recognizes periods, question marks, and exclamation marks, "
        "but protects decimal numbers, common abbreviations such as e.g. and Dr., person initials, and closing quotation marks. Each non-empty paragraph is processed separately, and remaining text without final punctuation still counts as one sentence. "
        "This approach is simple, transparent, and suitable for checking the system during a course demonstration.",
    )

    add_heading(document, "5. Results and User Interface")
    paragraph = document.add_paragraph()
    write_body(
        paragraph,
        "The interface highlights matched terms in the result snippet and shows the PMID, PMCID, and BM25 score for each document. "
        "Users can inspect term counts in the abstract and body, open the original record on PubMed, and review every detected sentence. "
        "The system also provides prefix autocomplete and a Did you mean suggestion based on Damerau-Levenshtein edit distance and term frequency. "
        "These features help users enter a query while keeping the retrieval process understandable rather than hidden.",
    )

    add_heading(document, "6. Implementation and Data Reliability")
    paragraph = document.add_paragraph()
    write_body(
        paragraph,
        "Helix is implemented in Python with Streamlit for the web interface. XML parsing uses standard XML tools, while NLTK provides the English stopword list and Porter stemmer. "
        "The application uses the official NCBI services when it retrieves a record by PMID or PMCID. Requests are rate limited, retried after temporary failures, and cached to avoid unnecessary repeated calls. "
        "Users may also upload XML directly, which is useful when a network connection is not available.",
    )
    paragraph = document.add_paragraph()
    write_body(
        paragraph,
        "The raw XML source can be kept locally, and the processed collection is stored as JSONL documents plus a saved index for fast loading. "
        "A small SQLite library records saved articles. Upload validation, content hashing, and atomic file writing help keep local data consistent when files are added or updated.",
    )

    add_heading(document, "7. Full Text Scope and Demonstration")
    paragraph = document.add_paragraph()
    write_body(
        paragraph,
        "The current demonstration collection mainly uses PubMed abstract XML, which is enough to show the retrieval and statistics pipeline. "
        "The same parser also supports PMC JATS XML. When a PMC record includes body text, Helix can index and search the abstract, keywords, and body together. "
        "The title remains excluded from matching so that the result is based on the article content. During a demonstration, a user can upload XML or retrieve a PMID, inspect the calculated statistics, run an AND, OR, or PHRASE query, and verify the highlighted terms and BM25 order. "
        "This provides a clear path from input data to visible retrieval evidence.",
    )

    conclusion = document.add_paragraph()
    conclusion.paragraph_format.space_before = Pt(5)
    write_body(
        conclusion,
        "In summary, Helix combines XML parsing, reproducible text rules, inverted indexing, Boolean matching, BM25 ranking, document statistics, and rule-based sentence detection in one application. "
        "It provides an inspectable search engine that directly addresses the retrieval, counting, and visualization goals of the assignment.",
    )

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
